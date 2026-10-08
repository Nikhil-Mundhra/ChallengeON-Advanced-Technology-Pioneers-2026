from __future__ import annotations

import json
import pickletools
import re
import subprocess
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any


class ToolError(ValueError):
    pass


class ReadOnlyTools:
    def __init__(self, root: Path, *, max_chars: int = 30_000):
        self.root = root.resolve()
        self.max_chars = max_chars

    def _path(self, raw: str) -> Path:
        candidate = (self.root / raw).resolve()
        try:
            candidate.relative_to(self.root)
        except ValueError as exc:
            raise ToolError("path escapes repository root") from exc
        return candidate

    def execute(self, name: str, args: dict[str, Any]) -> dict[str, Any]:
        methods = {
            "list_files": self.list_files,
            "read_text": self.read_text,
            "search_text": self.search_text,
            "file_info": self.file_info,
            "json_inspect": self.json_inspect,
            "tabular_profile": self.tabular_profile,
            "duckdb_query": self.duckdb_query,
            "git_inventory": self.git_inventory,
            "repository_inventory": self.repository_inventory,
            "parquet_metadata": self.parquet_metadata,
            "excel_metadata": self.excel_metadata,
            "pickle_metadata": self.pickle_metadata,
            "structured_artifact_inventory": self.structured_artifact_inventory,
        }
        if name not in methods:
            raise ToolError(f"unknown tool {name!r}")
        return methods[name](**args)

    def list_files(self, glob: str = "**/*", max_results: int = 200) -> dict[str, Any]:
        max_results = min(max(int(max_results), 1), 1000)
        values = []
        for path in sorted(self.root.glob(glob)):
            relative = path.relative_to(self.root)
            if path.is_file() and not ({".git", ".venv"} & set(relative.parts)):
                values.append(str(path.relative_to(self.root)))
                if len(values) >= max_results:
                    break
        return {"files": values, "truncated": len(values) == max_results}

    def read_text(
        self, path: str, start_line: int = 1, end_line: int = 300
    ) -> dict[str, Any]:
        file_path = self._path(path)
        if not file_path.is_file():
            raise ToolError(f"not a file: {path}")
        start_line = max(int(start_line), 1)
        end_line = min(max(int(end_line), start_line), start_line + 499)
        try:
            lines = file_path.read_text(encoding="utf-8").splitlines()
        except UnicodeDecodeError as exc:
            raise ToolError("file is binary; use file_info or a structured data tool") from exc
        selected = "\n".join(
            f"{number}: {lines[number - 1]}"
            for number in range(start_line, min(end_line, len(lines)) + 1)
        )
        return {
            "path": path,
            "start_line": start_line,
            "end_line": min(end_line, len(lines)),
            "total_lines": len(lines),
            "content": selected[: self.max_chars],
            "truncated": len(selected) > self.max_chars,
        }

    def search_text(
        self, query: str, glob: str = "**/*", max_results: int = 100
    ) -> dict[str, Any]:
        if not query or len(query) > 500:
            raise ToolError("query must contain 1-500 characters")
        pattern = re.compile(query, re.IGNORECASE)
        results: list[dict[str, Any]] = []
        for path in sorted(self.root.glob(glob)):
            if not path.is_file() or ".git" in path.parts or path.stat().st_size > 5_000_000:
                continue
            try:
                lines = path.read_text(encoding="utf-8").splitlines()
            except (UnicodeDecodeError, OSError):
                continue
            for number, line in enumerate(lines, 1):
                if pattern.search(line):
                    results.append(
                        {
                            "path": str(path.relative_to(self.root)),
                            "line": number,
                            "text": line[:500],
                        }
                    )
                    if len(results) >= min(int(max_results), 500):
                        return {"matches": results, "truncated": True}
        return {"matches": results, "truncated": False}

    def file_info(self, path: str) -> dict[str, Any]:
        file_path = self._path(path)
        stat = file_path.stat()
        return {
            "path": path,
            "is_file": file_path.is_file(),
            "is_directory": file_path.is_dir(),
            "size_bytes": stat.st_size,
            "modified_ns": stat.st_mtime_ns,
            "suffix": file_path.suffix,
        }

    def json_inspect(self, path: str, key: str | None = None) -> dict[str, Any]:
        file_path = self._path(path)
        value: Any = json.loads(file_path.read_text(encoding="utf-8"))
        if key:
            for part in key.split("."):
                if isinstance(value, dict):
                    value = value[part]
                elif isinstance(value, list):
                    value = value[int(part)]
                else:
                    raise ToolError(f"cannot descend through {part!r}")
        rendered = json.dumps(value, indent=2, ensure_ascii=False)
        summary = {
            "type": type(value).__name__,
            "length": len(value) if isinstance(value, (dict, list)) else None,
            "value": rendered[: self.max_chars],
            "truncated": len(rendered) > self.max_chars,
        }
        return summary

    def tabular_profile(
        self, path: str, sheet: str | None = None, max_categories: int = 20
    ) -> dict[str, Any]:
        import pandas as pd

        file_path = self._path(path)
        suffix = file_path.suffix.lower()
        if suffix in {".xlsx", ".xls"}:
            source = pd.read_excel(file_path, sheet_name=sheet if sheet else 0)
        elif suffix == ".parquet":
            source = pd.read_parquet(file_path)
        elif suffix in {".csv", ".tsv"}:
            source = pd.read_csv(file_path, sep="\t" if suffix == ".tsv" else ",")
        else:
            raise ToolError(f"unsupported table format: {suffix}")
        profile: dict[str, Any] = {
            "path": path,
            "sheet": sheet,
            "rows": len(source),
            "columns": len(source.columns),
            "duplicate_rows": int(source.duplicated().sum()),
            "fields": [],
        }
        for column in source.columns:
            series = source[column]
            field: dict[str, Any] = {
                "name": str(column),
                "dtype": str(series.dtype),
                "nulls": int(series.isna().sum()),
                "distinct_non_null": int(series.nunique(dropna=True)),
            }
            non_null = series.dropna()
            if len(non_null) and (
                pd.api.types.is_numeric_dtype(series) or pd.api.types.is_datetime64_any_dtype(series)
            ):
                field["min"] = str(non_null.min())
                field["max"] = str(non_null.max())
            elif 0 < field["distinct_non_null"] <= int(max_categories):
                field["values"] = [str(value) for value in non_null.unique()[:max_categories]]
            profile["fields"].append(field)
        rendered = json.dumps(profile, ensure_ascii=False)
        if len(rendered) > self.max_chars:
            profile["fields"] = profile["fields"][:50]
            profile["truncated"] = True
        return profile

    def duckdb_query(
        self, database: str, sql: str, max_rows: int = 200
    ) -> dict[str, Any]:
        import duckdb

        normalized = re.sub(r"\s+", " ", sql.strip()).lower()
        if not re.match(r"^(select|with|describe|show)\b", normalized):
            raise ToolError("only SELECT, WITH, DESCRIBE, or SHOW queries are allowed")
        forbidden = re.compile(
            r"\b(copy|attach|detach|install|load|call|export|import|create|replace|"
            r"update|delete|insert|drop|alter|truncate|vacuum)\b",
            re.I,
        )
        if forbidden.search(sql) or ";" in sql.rstrip(";"):
            raise ToolError("query contains a forbidden or multi-statement operation")
        database_path = self._path(database)
        connection = duckdb.connect(str(database_path), read_only=True)
        try:
            cursor = connection.execute(sql)
            columns = [item[0] for item in cursor.description]
            rows = cursor.fetchmany(min(max(int(max_rows), 1), 1000) + 1)
        finally:
            connection.close()
        limit = min(max(int(max_rows), 1), 1000)
        return {
            "columns": columns,
            "rows": [[self._json_value(value) for value in row] for row in rows[:limit]],
            "truncated": len(rows) > limit,
        }

    def git_inventory(self, operation: str = "status") -> dict[str, Any]:
        commands = {
            "status": ["git", "status", "--short"],
            "tracked": ["git", "ls-files"],
        }
        if operation not in commands:
            raise ToolError("operation must be status or tracked")
        result = subprocess.run(
            commands[operation],
            cwd=self.root,
            check=False,
            capture_output=True,
            text=True,
            timeout=30,
        )
        return {
            "operation": operation,
            "exit_code": result.returncode,
            "output": result.stdout[: self.max_chars],
            "stderr": result.stderr[:2000],
        }

    def repository_inventory(self, max_files: int = 2000) -> dict[str, Any]:
        """Return classified file metadata without flooding the model with runtime packages."""
        max_files = min(max(int(max_files), 1), 5000)
        tracked_result = subprocess.run(
            ["git", "ls-files", "-z"],
            cwd=self.root,
            check=False,
            capture_output=True,
            timeout=30,
        )
        tracked = {
            value.decode("utf-8", "replace")
            for value in tracked_result.stdout.split(b"\0")
            if value
        }
        status_result = subprocess.run(
            ["git", "status", "--short", "-z", "--untracked-files=all"],
            cwd=self.root,
            check=False,
            capture_output=True,
            timeout=30,
        )
        status: dict[str, str] = {}
        for entry in status_result.stdout.split(b"\0"):
            if len(entry) >= 4:
                decoded = entry.decode("utf-8", "replace")
                status[decoded[3:]] = decoded[:2]

        records = []
        excluded_counts: Counter[str] = Counter()
        for path in sorted(self.root.rglob("*")):
            if not path.is_file():
                continue
            relative = path.relative_to(self.root)
            if ".git" in relative.parts:
                excluded_counts[".git internals"] += 1
                continue
            if ".venv" in relative.parts:
                excluded_counts["virtual-environment packages"] += 1
                continue
            stat = path.stat()
            relative_text = str(relative)
            records.append(
                {
                    "path": relative_text,
                    "size": stat.st_size,
                    "mtime": datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(),
                    "extension": path.suffix.lower() or "<none>",
                    "tracked": relative_text in tracked,
                    "git_status": status.get(relative_text, ""),
                    "role": self._classify(relative),
                }
            )
            if len(records) >= max_files:
                break
        roles = Counter(record["role"] for record in records)
        untracked = [record["path"] for record in records if record["git_status"] == "??"]
        unknown = [record["path"] for record in records if record["role"] == "unknown"]
        return {
            "files": records,
            "file_count": len(records),
            "role_counts": dict(sorted(roles.items())),
            "untracked": untracked,
            "unknown": unknown,
            "excluded_counts": dict(excluded_counts),
            "truncated": len(records) >= max_files,
        }

    def parquet_metadata(self, path: str) -> dict[str, Any]:
        import pyarrow.parquet as parquet

        file_path = self._path(path)
        source = parquet.ParquetFile(file_path)
        metadata = source.metadata
        row_groups = []
        for index in range(metadata.num_row_groups):
            group = metadata.row_group(index)
            row_groups.append(
                {
                    "index": index,
                    "rows": group.num_rows,
                    "total_byte_size": group.total_byte_size,
                    "columns": group.num_columns,
                }
            )
        return {
            "path": path,
            "rows": metadata.num_rows,
            "row_groups": row_groups,
            "created_by": metadata.created_by,
            "schema": str(source.schema_arrow),
            "metadata": {
                key.decode("utf-8", "replace"): value.decode("utf-8", "replace")[:2000]
                for key, value in (metadata.metadata or {}).items()
            },
        }

    def excel_metadata(self, path: str) -> dict[str, Any]:
        from openpyxl import load_workbook

        file_path = self._path(path)
        workbook = load_workbook(file_path, read_only=False, data_only=False)
        try:
            sheets = []
            for worksheet in workbook.worksheets:
                hidden_rows = sum(1 for value in worksheet.row_dimensions.values() if value.hidden)
                hidden_columns = sum(1 for value in worksheet.column_dimensions.values() if value.hidden)
                formulas = 0
                for row in worksheet.iter_rows():
                    formulas += sum(
                        1 for cell in row if isinstance(cell.value, str) and cell.value.startswith("=")
                    )
                sheets.append(
                    {
                        "title": worksheet.title,
                        "state": worksheet.sheet_state,
                        "max_row": worksheet.max_row,
                        "max_column": worksheet.max_column,
                        "merged_ranges": [str(value) for value in worksheet.merged_cells.ranges],
                        "hidden_rows": hidden_rows,
                        "hidden_columns": hidden_columns,
                        "formula_cells": formulas,
                        "auto_filter": str(worksheet.auto_filter.ref or ""),
                    }
                )
            return {
                "path": path,
                "sheets": sheets,
                "defined_names": sorted(str(value) for value in workbook.defined_names),
            }
        finally:
            workbook.close()

    def pickle_metadata(self, path: str, max_globals: int = 200) -> dict[str, Any]:
        """Inspect pickle opcodes without executing the pickle."""
        file_path = self._path(path)
        data = file_path.read_bytes()
        protocols: set[int] = set()
        globals_seen: list[str] = []
        opcode_counts: Counter[str] = Counter()
        for opcode, argument, _ in pickletools.genops(data):
            opcode_counts[opcode.name] += 1
            if opcode.name == "PROTO" and isinstance(argument, int):
                protocols.add(argument)
            if opcode.name in {"GLOBAL", "STACK_GLOBAL"} and len(globals_seen) < max_globals:
                globals_seen.append(str(argument))
        return {
            "path": path,
            "size_bytes": len(data),
            "protocols": sorted(protocols),
            "globals": globals_seen,
            "opcode_counts": dict(sorted(opcode_counts.items())),
            "executed": False,
            "note": "Opcode inspection cannot prove the runtime top-level object type without unpickling.",
        }

    def structured_artifact_inventory(
        self, directory: str = "lake", database: str = "lake/analytics.duckdb"
    ) -> dict[str, Any]:
        """Inventory all structured lake artifacts using safe metadata operations."""
        import duckdb

        directory_path = self._path(directory)
        database_path = self._path(database)
        result: dict[str, Any] = {"database": {}, "artifacts": []}
        connection = duckdb.connect(str(database_path), read_only=True)
        try:
            objects = connection.execute(
                """
                SELECT table_schema, table_name, table_type
                FROM information_schema.tables
                WHERE table_schema NOT IN ('information_schema', 'pg_catalog')
                ORDER BY table_schema, table_name
                """
            ).fetchall()
            columns = connection.execute(
                """
                SELECT table_schema, table_name, column_name, data_type, is_nullable,
                       ordinal_position
                FROM information_schema.columns
                WHERE table_schema NOT IN ('information_schema', 'pg_catalog')
                ORDER BY table_schema, table_name, ordinal_position
                """
            ).fetchall()
            columns_by_object: dict[tuple[str, str], list[dict[str, Any]]] = {}
            for schema, table, column, data_type, nullable, position in columns:
                columns_by_object.setdefault((schema, table), []).append(
                    {
                        "name": column,
                        "type": data_type,
                        "nullable": nullable,
                        "position": position,
                    }
                )
            database_objects = []
            for schema, name, kind in objects:
                escaped = str(name).replace('"', '""')
                count = connection.execute(f'SELECT count(*) FROM "{escaped}"').fetchone()[0]
                database_objects.append(
                    {
                        "schema": schema,
                        "name": name,
                        "type": kind,
                        "rows": count,
                        "columns": [
                            f"{column['name']}:{column['type']}"
                            for column in columns_by_object.get((schema, name), [])
                        ],
                    }
                )
            views = connection.execute(
                "SELECT schema_name, view_name, sql FROM duckdb_views() "
                "WHERE internal = false ORDER BY schema_name, view_name"
            ).fetchall()
            result["database"] = {
                "path": database,
                "objects": database_objects,
                "view_definitions": [
                    {
                        "schema": schema,
                        "name": name,
                        "sql": sql[:1000],
                        "sql_truncated": len(sql) > 1000,
                    }
                    for schema, name, sql in views
                ],
            }
        finally:
            connection.close()

        for path in sorted(directory_path.rglob("*")):
            if not path.is_file() or path.resolve() == database_path:
                continue
            relative = str(path.relative_to(self.root))
            suffix = path.suffix.lower()
            try:
                if suffix == ".parquet":
                    metadata = self.parquet_metadata(relative)
                    metadata.pop("metadata", None)
                    result["artifacts"].append({"kind": "parquet", **metadata})
                elif suffix == ".json":
                    value = json.loads(path.read_text(encoding="utf-8"))
                    result["artifacts"].append(
                        {
                            "kind": "json",
                            "path": relative,
                            "top_level_type": type(value).__name__,
                            "top_level_length": len(value) if isinstance(value, (dict, list)) else None,
                            "key_tree": self._json_key_tree(value),
                        }
                    )
                elif suffix in {".pkl", ".pickle"}:
                    result["artifacts"].append(
                        {"kind": "pickle", **self.pickle_metadata(relative)}
                    )
            except Exception as exc:
                result["artifacts"].append(
                    {"kind": suffix.lstrip("."), "path": relative, "error": str(exc)}
                )
        return result

    @classmethod
    def _json_key_tree(cls, value: Any, depth: int = 0) -> Any:
        if depth >= 2:
            return f"<{type(value).__name__}>"
        if isinstance(value, dict):
            entries = {
                str(key): cls._json_key_tree(child, depth + 1)
                for key, child in list(value.items())[:30]
            }
            if len(value) > 30:
                entries["<truncated_keys>"] = len(value) - 30
            return entries
        if isinstance(value, list):
            return {
                "length": len(value),
                "item_example": cls._json_key_tree(value[0], depth + 1) if value else None,
            }
        return f"<{type(value).__name__}>"

    @staticmethod
    def _classify(path: Path) -> str:
        text = str(path).lower()
        suffix = path.suffix.lower()
        if "01a - dct dataset" in text:
            return "raw source"
        if "__pycache__" in path.parts or ".pytest_cache" in path.parts or suffix == ".pyc":
            return "cache"
        if text.startswith("lake/curated/"):
            return "model artifact" if suffix in {".pkl", ".json"} else "curated artifact"
        if text == "lake/analytics.duckdb":
            return "database"
        if text.startswith("output/") or text.startswith("tmp/pdfs/"):
            return "report" if suffix in {".pdf", ".png"} else "generated output"
        if suffix in {".py", ".sql", ".html", ".css", ".js"}:
            return "source code"
        if suffix in {".md", ".pdf"}:
            return "documentation"
        if path.name in {"Makefile", "pyproject.toml", ".gitignore"}:
            return "configuration"
        if text.startswith("audit/"):
            return "audit state"
        return "unknown"

    @staticmethod
    def _json_value(value: Any) -> Any:
        if value is None or isinstance(value, (str, int, float, bool)):
            return value
        return str(value)


TOOL_DOCUMENTATION = """
Available read-only tools. Request exactly one tool at a time.

- list_files: {"glob":"pattern", "max_results":200}
- read_text: {"path":"relative/path", "start_line":1, "end_line":300}
- search_text: {"query":"regular expression", "glob":"**/*.py", "max_results":100}
- file_info: {"path":"relative/path"}
- json_inspect: {"path":"relative.json", "key":"optional.dotted.path"}
- tabular_profile: {"path":"file.xlsx|parquet|csv", "sheet":"optional"}
- duckdb_query: {"database":"lake/analytics.duckdb", "sql":"SELECT ...", "max_rows":200}
- git_inventory: {"operation":"status|tracked"}
- repository_inventory: {"max_files":2000}; returns path, size, mtime, Git state,
  inferred role, exclusions, and exception lists while excluding .git internals and
  virtual-environment packages as permitted by the checklist
- parquet_metadata: {"path":"file.parquet"}; schema, rows, row groups, writer metadata
- excel_metadata: {"path":"file.xlsx"}; sheet dimensions, visibility, formulas,
  merged ranges, filters, and defined names
- pickle_metadata: {"path":"file.pkl"}; safe opcode/protocol/global inspection without execution
- structured_artifact_inventory: {"directory":"lake", "database":"lake/analytics.duckdb"};
  inventories all DuckDB objects and counts, Parquet schemas and row groups, JSON key
  trees, and pickle opcodes in one bounded read-only call
""".strip()

"""Slide primitives on python-pptx. Every shape gets a stable `name` ("title", "hero.value",
"col2.head", "bg.grid", ...) so people can find it in PowerPoint and tests can assert on it.

Text boxes have zero insets, no autofit and exact line spacing, so a box's height is its measured
line count times the style's line height (typeset.py); text that does not fit raises DeckOverflow.
Shapes never take a theme style (no effectRef), so no renderer adds a shadow.
"""

from __future__ import annotations

from io import BytesIO
from typing import Optional, Sequence

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from tourism_twin.reporting.deck import grid, theme, typeset
from tourism_twin.reporting.deck.grid import Box
from tourism_twin.reporting.deck.typeset import DeckOverflow, Style

ALIGN = {"left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT, "center": PP_ALIGN.CENTER}


def new_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(theme.SLIDE_W), Inches(theme.SLIDE_H)
    return prs


def new_slide(prs: Presentation, tone: str):
    """A blank slide with the tone's background colour ("light" or "dark")."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _rgb(typeset.PALETTES[tone]["background"])
    slide.tone = tone  # read by the primitives below
    return slide


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def _color(slide, role: str) -> RGBColor:
    return _rgb(typeset.PALETTES[slide.tone][role])


def _emu(box: Box):
    return Inches(box.x), Inches(box.y), Inches(box.w), Inches(box.h)


def _no_style(shape) -> None:
    """Drop the theme style reference python-pptx adds to autoshapes (its effectRef draws a shadow
    in LibreOffice)."""
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


def text(slide, name: str, box: Box, content: str, style: Style | str, anchor: str = "top") -> float:
    """Draw `content` (markup allowed) in `box`; returns the height the text uses. Raises
    DeckOverflow when the wrapped lines need more height than the box has."""
    style = typeset.STYLES[style] if isinstance(style, str) else style
    lines = typeset.wrap(content, style, box.w)
    used = len(lines) * style.line
    if used > box.h + 1e-6:
        raise DeckOverflow(f"{name}: {len(lines)} lines need {used:.2f} in, the box has {box.h:.2f} in: "
                           f"shorten {typeset.plain(content)[:60]!r}")
    shape = slide.shapes.add_textbox(*_emu(box))
    shape.name = name
    frame = shape.text_frame
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
    for i, line in enumerate(content.split("\n")):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.alignment = ALIGN[style.align]
        paragraph.line_spacing = Pt(style.size * style.leading)
        paragraph.space_before = paragraph.space_after = Pt(0)
        for run in typeset.parse(line, style):
            r = paragraph.add_run()
            r.text = run.text
            r.font.name = theme.FONT
            r.font.size = Pt(style.size)
            r.font.bold = run.weight == "bold"
            r.font.color.rgb = _color(slide, run.role)
            if run.weight == "semibold":
                r.font.name = f"{theme.FONT} SemiBold"
    return used


def fit(content: str, style: Style | str, w: float) -> float:
    """Height `content` needs in a box `w` wide."""
    style = typeset.STYLES[style] if isinstance(style, str) else style
    return typeset.height(content, style, w)


def rule(slide, name: str, x: float, y: float, length: float, vertical: bool = False, role: str = "hairline") -> None:
    """A hairline: a thin filled rectangle, so it renders identically everywhere."""
    thickness = theme.HAIRLINE_PT / 72
    box = Box(x, y, thickness, length) if vertical else Box(x, y, length, thickness)
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, *_emu(box))
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = _color(slide, role)
    shape.line.fill.background()
    _no_style(shape)


def fill(slide, name: str, box: Box, role: str = "background"):
    """A borderless filled rectangle (the full-bleed background: some viewers ignore the slide
    background fill, a shape renders everywhere)."""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, *_emu(box))
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = _color(slide, role)
    shape.line.fill.background()
    _no_style(shape)
    return shape


def card(slide, name: str, box: Box) -> None:
    """White card with a 1 pt hairline border, square corners, no shadow."""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, *_emu(box))
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = _color(slide, "card")
    shape.line.color.rgb = _color(slide, "hairline")
    shape.line.width = Pt(theme.HAIRLINE_PT)
    _no_style(shape)


def arrow(slide, name: str, start: tuple[float, float], end: tuple[float, float]) -> None:
    """Accent connector with an arrowhead at `end`."""
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(start[0]), Inches(start[1]),
                                      Inches(end[0]), Inches(end[1]))
    line.name = name
    line.line.color.rgb = _color(slide, "accent")
    line.line.width = Pt(theme.CONNECTOR_PT)
    ln = line.line._get_or_add_ln()
    etree.SubElement(ln, qn("a:tailEnd"), type="triangle", w="med", len="med")
    _no_style(line)


def picture(slide, name: str, png: bytes, box: Box, fit_inside: bool = True):
    """Place PNG bytes in `box` (keeping the aspect ratio, top-left aligned) and return the shape."""
    from PIL import Image

    with Image.open(BytesIO(png)) as image:
        ratio = image.width / image.height
    w, h = box.w, box.h
    if fit_inside:
        w, h = (box.w, box.w / ratio) if ratio > box.w / box.h else (box.h * ratio, box.h)
    shape = slide.shapes.add_picture(BytesIO(png), Inches(box.x), Inches(box.y), Inches(w), Inches(h))
    shape.name = name
    return shape


def send_to_back(slide, shapes: Sequence) -> None:
    tree = slide.shapes._spTree
    for offset, shape in enumerate(shapes):
        tree.remove(shape._element)
        tree.insert(2 + offset, shape._element)


def notes(slide, content: str) -> None:
    if content:
        slide.notes_slide.notes_text_frame.text = content


def chrome(slide, kicker: Optional[str], footer: str, page: int, total: int, icon: Optional[bytes] = None) -> None:
    """Kicker (with the plane icon in the margin), footer hairline, footer text and page number."""
    if kicker:
        k, _, _ = grid.header(1, 1)
        text(slide, "kicker", k, kicker, "kicker")
        if icon is not None:
            size = theme.ICON
            picture(slide, "kicker.icon", icon,
                    Box(theme.MARGIN - size - theme.GUTTER, k.y + (k.h - size) / 2, size, size))
    rule(slide, "footer.rule", theme.MARGIN, theme.FOOTER_RULE_Y, grid.CONTENT_W)
    page_w = grid.COLUMN_W * 2
    foot = typeset.STYLES["footer"]
    text(slide, "footer", Box(theme.MARGIN, theme.FOOTER_TEXT_Y, grid.CONTENT_W - page_w, foot.line), footer, foot)
    text(slide, "page", Box(theme.MARGIN + grid.CONTENT_W - page_w, theme.FOOTER_TEXT_Y, page_w, foot.line),
         f"{page} / {total}", "page")


__all__ = ["DeckOverflow", "new_presentation", "new_slide", "text", "fit", "rule", "card", "arrow", "picture",
           "send_to_back", "notes", "chrome", "fill"]

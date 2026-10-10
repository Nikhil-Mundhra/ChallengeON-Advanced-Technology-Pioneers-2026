"""Slide primitives on python-pptx: one 16:9 layout family in the repo palette.

Every slide has the same frame: a title, a one-line message under it, a body (bullet tree,
figure, table, or two of these side by side) and a footer with the deck name and page number.
Bullet trees encode hierarchy by indent: a parent line, then its children one level in.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from tourism_twin.reporting.palette import BLUE, INK, LINE, MUTED, NAVY, PALE, SKY

WIDTH, HEIGHT = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.6)
BODY_TOP = Inches(1.75)
BODY_BOTTOM = HEIGHT - Inches(0.6)
FONT = "Calibri"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#"))


def new_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = WIDTH, HEIGHT
    return prs


def _text(frame, text: str, size: float, color: str, bold: bool = False, align=PP_ALIGN.LEFT) -> None:
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.size, run.font.bold, run.font.name = Pt(size), bold, FONT
    run.font.color.rgb = rgb(color)


def frame_slide(prs: Presentation, title: str, message: str, page: int, total: int, footer: str):
    """A blank slide with the shared frame; returns the slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    band = slide.shapes.add_shape(1, 0, 0, WIDTH, Inches(0.12))
    band.fill.solid()
    band.fill.fore_color.rgb = rgb(BLUE)
    band.line.fill.background()

    title_box = slide.shapes.add_textbox(MARGIN, Inches(0.35), WIDTH - 2 * MARGIN, Inches(0.7))
    _text(title_box.text_frame, title, 30, NAVY, bold=True)
    if message:
        message_box = slide.shapes.add_textbox(MARGIN, Inches(1.0), WIDTH - 2 * MARGIN, Inches(0.6))
        _text(message_box.text_frame, message, 17, BLUE)

    rule = slide.shapes.add_connector(1, MARGIN, Inches(1.6), WIDTH - MARGIN, Inches(1.6))
    rule.line.color.rgb = rgb(LINE)

    foot = slide.shapes.add_textbox(MARGIN, HEIGHT - Inches(0.45), WIDTH - 2 * MARGIN, Inches(0.3))
    _text(foot.text_frame, footer, 10, MUTED)
    number = slide.shapes.add_textbox(WIDTH - MARGIN - Inches(1.0), HEIGHT - Inches(0.45), Inches(1.0), Inches(0.3))
    _text(number.text_frame, f"{page} / {total}", 10, MUTED, align=PP_ALIGN.RIGHT)
    return slide


def bullet_tree(slide, items: Sequence[Dict[str, Any] | str], left: Emu, top: Emu, width: Emu, height: Emu) -> None:
    """Parent lines bold in navy; children one indent level in, ink; grandchildren muted."""
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    first = True

    def add(text: str, level: int) -> None:
        nonlocal first
        paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
        first = False
        paragraph.level = min(level, 2)
        paragraph.space_before = Pt(10 if level == 0 else 3)
        marker = "" if level == 0 else ("–  " if level == 1 else "·  ")
        run = paragraph.add_run()
        run.text = marker + text
        run.font.name = FONT
        run.font.size = Pt((19, 16, 14)[min(level, 2)])
        run.font.bold = level == 0
        run.font.color.rgb = rgb((NAVY, INK, MUTED)[min(level, 2)])

    def walk(nodes: Sequence[Dict[str, Any] | str], level: int) -> None:
        for node in nodes:
            if isinstance(node, str):
                add(node, level)
            else:
                add(node["text"], level)
                walk(node.get("children", []), level + 1)

    walk(items, 0)


def picture(slide, path: Path, left: Emu, top: Emu, width: Emu, height: Emu) -> None:
    """Fit the image inside the box, keeping its aspect ratio, centred."""
    from PIL import Image

    with Image.open(path) as image:
        ratio = image.width / image.height
    box_ratio = width / height
    w, h = (width, int(width / ratio)) if ratio > box_ratio else (int(height * ratio), height)
    slide.shapes.add_picture(str(path), left + (width - w) // 2, top + (height - h) // 2, w, h)


def table(slide, header: List[str], rows: List[List[str]], left: Emu, top: Emu, width: Emu,
          highlight_row: Optional[int] = None) -> None:
    shape = slide.shapes.add_table(len(rows) + 1, len(header), left, top, width, Inches(0.42) * (len(rows) + 1))
    grid = shape.table
    for col, label in enumerate(header):
        cell = grid.cell(0, col)
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(NAVY)
        cell.text_frame.text = ""
        _text(cell.text_frame, label, 14, "#FFFFFF", bold=True)
    for r, row in enumerate(rows, start=1):
        for col, value in enumerate(row):
            cell = grid.cell(r, col)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(SKY if r - 1 == highlight_row else (PALE if r % 2 else "#FFFFFF"))
            cell.text_frame.text = ""
            _text(cell.text_frame, value, 14, NAVY if r - 1 == highlight_row else INK, bold=r - 1 == highlight_row)


def notes(slide, text: str) -> None:
    if text:
        slide.notes_slide.notes_text_frame.text = text


def body_box(split: Optional[str] = None):
    """Body regions: full width, or (left, right) for 'left' text / 'right' figure layouts."""
    width = WIDTH - 2 * MARGIN
    height = BODY_BOTTOM - BODY_TOP
    if split is None:
        return (MARGIN, BODY_TOP, width, height)
    left_w = int(width * 0.46)
    gap = Inches(0.3)
    return (MARGIN, BODY_TOP, left_w, height), (MARGIN + left_w + gap, BODY_TOP, width - left_w - gap, height)
